# Carga cada formato de la base de conocimiento y lo divide en fragmentos (chunks)
# según su estructura. Cada fragmento lleva metadatos: fuente, tipo, audiencia, sección.
import csv
import json
import re
from pathlib import Path

from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_core.documents import Document
from langchain_text_splitters import MarkdownHeaderTextSplitter, RecursiveCharacterTextSplitter
from pypdf import PdfReader

BASE = Path(__file__).parent / "base_conocimiento"

TITULOS = {
    "politica_devoluciones.pdf": "Política de Devoluciones",
    "catalogo_productos.csv": "Catálogo de productos",
    "faq.json": "Preguntas frecuentes",
    "politica_envios.md": "Política de Envíos",
    "politica_garantias_sostenibilidad.docx": "Garantías y Sostenibilidad",
    "guia_atencion_interna.txt": "Guía interna de atención",
}

# Una sección más larga que esto se parte con el splitter recursivo.
MAX_CHARS = 1200
recursivo = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=150)


def _chunks(cuerpo, ruta, audiencia, seccion, con_encabezado):
    """Convierte el texto de una sección en uno o varios Document."""
    cuerpo = cuerpo.strip()
    if len(cuerpo) < 40:
        return []
    partes = [cuerpo] if len(cuerpo) <= MAX_CHARS else recursivo.split_text(cuerpo)
    docs = []
    for parte in partes:
        # Encabezado contextual: documento + sección. Un fragmento aislado
        # ("5 a 7 días hábiles") queda más claro si dice de dónde viene.
        texto = f"{TITULOS[ruta.name]} | {seccion}\n{parte}" if con_encabezado else parte
        docs.append(Document(page_content=texto, metadata={
            "fuente": ruta.name, "tipo": ruta.suffix[1:], "audiencia": audiencia, "seccion": seccion}))
    return docs


# --- PDF: por secciones numeradas ("1. Alcance", "2. Plazo...") -------------
def _texto_pdf(ruta):
    return "\n".join(p.extract_text() for p in PdfReader(ruta).pages)


def cargar_pdf(ruta, audiencia, con_encabezado=True):
    secciones, titulo, cuerpo, esperado = [], "Encabezado", [], 1
    for linea in _texto_pdf(ruta).splitlines():
        m = re.match(r"^(\d+)\. ", linea.strip())
        # Es título si el número es el siguiente esperado, es corto y no termina en punto
        # (los pasos numerados de una lista sí terminan en punto).
        if m and int(m.group(1)) == esperado and len(linea) < 60 and not linea.strip().endswith("."):
            secciones.append((titulo, cuerpo))
            titulo, cuerpo, esperado = linea.strip(), [], esperado + 1
        else:
            cuerpo.append(linea)
    secciones.append((titulo, cuerpo))
    docs = []
    for titulo, lineas in secciones:
        if titulo != "Encabezado":
            docs += _chunks("\n".join(lineas), ruta, audiencia, titulo, con_encabezado)
    return docs


# --- Markdown: por encabezados (##, ###) ------------------------------------
def cargar_md(ruta, audiencia, con_encabezado=True):
    splitter = MarkdownHeaderTextSplitter(
        [("#", "h1"), ("##", "h2"), ("###", "h3")], strip_headers=False)
    docs = []
    for d in splitter.split_text(ruta.read_text(encoding="utf-8")):
        seccion = " > ".join(d.metadata[k] for k in ("h2", "h3") if k in d.metadata) or "Introducción"
        docs += _chunks(d.page_content, ruta, audiencia, seccion, con_encabezado)
    return docs


# --- DOCX: por títulos (Heading 1); las tablas se convierten en una línea por fila ---
def _bloques_docx(ruta):
    d = DocxDocument(ruta)
    for hijo in d.element.body.iterchildren():
        if hijo.tag.endswith("}p"):
            p = Paragraph(hijo, d)
            if p.text.strip():
                yield p.style.name, p.text.strip()
        elif hijo.tag.endswith("}tbl"):
            t = Table(hijo, d)
            encabezados = [c.text.strip() for c in t.rows[0].cells]
            for fila in t.rows[1:]:
                yield "Tabla", "; ".join(f"{h}: {c.text.strip()}" for h, c in zip(encabezados, fila.cells))


def cargar_docx(ruta, audiencia, con_encabezado=True):
    secciones, titulo, cuerpo = [], "Introducción", []
    for estilo, texto in _bloques_docx(ruta):
        if estilo == "Heading 1":
            secciones.append((titulo, cuerpo))
            titulo, cuerpo = texto, []
        elif estilo != "Title":
            cuerpo.append("- " + texto if estilo.startswith("List") else texto)
    secciones.append((titulo, cuerpo))
    docs = []
    for titulo, lineas in secciones:
        docs += _chunks("\n".join(lineas), ruta, audiencia, titulo, con_encabezado)
    return docs


# --- JSON (FAQ): un fragmento por pregunta con su respuesta -----------------
def cargar_json(ruta, audiencia, con_encabezado=True):
    docs = []
    for p in json.loads(ruta.read_text(encoding="utf-8"))["preguntas"]:
        cuerpo = f"Pregunta: {p['pregunta']}\nRespuesta: {p['respuesta']}"
        docs += _chunks(cuerpo, ruta, audiencia, p["categoria"], con_encabezado)
    return docs


# --- CSV (catálogo): un fragmento por producto, escrito como frase ----------
def cargar_csv(ruta, audiencia, con_encabezado=True):
    docs = []
    with open(ruta, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            precio = f"${int(r['precio_cop']):,}".replace(",", ".")
            partes = [
                f"{r['nombre']} ({r['sku']}). Categoría: {r['categoria']}. {r['descripcion']}",
                f"Material: {r['material']}. Origen: {r['origen']}. Precio: {precio} COP.",
                f"Disponibilidad: {r['disponibilidad']} (stock: {r['stock']}).",
                f"Devolución permitida: {r['devolucion_permitida']}."
                + (f" Motivo: {r['motivo_restriccion']}" if r["motivo_restriccion"] else ""),
                f"Garantía: {r['garantia_meses']} meses." if r["garantia_meses"] != "0" else "Sin garantía por ser consumible.",
            ]
            if r["certificaciones"]:
                partes.append(f"Certificaciones: {r['certificaciones']}.")
            docs += _chunks("\n".join(partes), ruta, audiencia, r["nombre"], con_encabezado)
    return docs


# --- TXT (guía interna): por secciones separadas con líneas de "=" ----------
def cargar_txt(ruta, audiencia, con_encabezado=True):
    partes = re.split(r"\n={20,}\n", ruta.read_text(encoding="utf-8"))
    docs = _chunks(partes[0], ruta, audiencia, "Introducción", con_encabezado)
    for titulo, cuerpo in zip(partes[1::2], partes[2::2]):
        docs += _chunks(cuerpo, ruta, audiencia, titulo.strip(), con_encabezado)
    return docs


CARGADORES = {".pdf": cargar_pdf, ".md": cargar_md, ".docx": cargar_docx,
              ".json": cargar_json, ".csv": cargar_csv, ".txt": cargar_txt}


def _archivos():
    for audiencia in ("publico", "interno"):
        for ruta in sorted((BASE / audiencia).iterdir()):
            if ruta.suffix in CARGADORES:
                yield ruta, audiencia


def cargar_base(con_encabezado=True):
    """Chunking por estructura: cada formato se divide según cómo está organizado."""
    docs = []
    for ruta, audiencia in _archivos():
        docs += CARGADORES[ruta.suffix](ruta, audiencia, con_encabezado)
    return docs


# --- Estrategias ingenuas, solo para compararlas en los experimentos ---------
def texto_plano(ruta):
    """El texto del archivo sin ninguna noción de estructura."""
    if ruta.suffix == ".pdf":
        return _texto_pdf(ruta)
    if ruta.suffix == ".docx":
        return "\n".join(t for _, t in _bloques_docx(ruta))
    if ruta.suffix == ".json":
        return "\n".join(f"{p['pregunta']} {p['respuesta']}"
                         for p in json.loads(ruta.read_text(encoding="utf-8"))["preguntas"])
    if ruta.suffix == ".csv":
        with open(ruta, encoding="utf-8") as f:
            return "\n".join("\n".join(f"{k}: {v}" for k, v in r.items()) + "\n" for r in csv.DictReader(f))
    return ruta.read_text(encoding="utf-8")


def cargar_tamano_fijo(palabras=200, solape=50):
    """Ventanas de N palabras con solapamiento (la estrategia de la demo del profesor)."""
    docs = []
    for ruta, audiencia in _archivos():
        w = texto_plano(ruta).split()
        for i in range(0, max(len(w) - solape, 1), palabras - solape):
            docs.append(Document(page_content=" ".join(w[i:i + palabras]), metadata={
                "fuente": ruta.name, "tipo": ruta.suffix[1:], "audiencia": audiencia, "seccion": ""}))
    return docs


def cargar_recursivo(tamano=1000, solape=200):
    """Splitter recursivo por caracteres (la estrategia de rag_ejemplo.py)."""
    sp = RecursiveCharacterTextSplitter(chunk_size=tamano, chunk_overlap=solape)
    docs = []
    for ruta, audiencia in _archivos():
        for parte in sp.split_text(texto_plano(ruta)):
            docs.append(Document(page_content=parte, metadata={
                "fuente": ruta.name, "tipo": ruta.suffix[1:], "audiencia": audiencia, "seccion": ""}))
    return docs
