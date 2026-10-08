# Genera los documentos de la base de conocimiento que no son texto plano:
#   - politica_devoluciones.pdf
#   - politica_garantias_sostenibilidad.docx
#   - catalogo_productos.csv
# EcoMarket es una empresa ficticia, así que este script emula los archivos que una
# empresa real tendría. Los demás documentos (md, json, txt) están escritos a mano.
#
# Uso (desde taller-2/):  python utilidades/generar_documentos.py
import csv
from pathlib import Path

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

DESTINO = Path(__file__).resolve().parent.parent / "base_conocimiento" / "publico"
DESTINO.mkdir(parents=True, exist_ok=True)


# ---------------------------------------------------------------------------
# 1. PDF: política de devoluciones
# ---------------------------------------------------------------------------
def generar_pdf():
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle("t", parent=estilos["Title"], fontSize=18, spaceAfter=4)
    sub = ParagraphStyle("s", parent=estilos["Normal"], textColor=colors.grey, spaceAfter=14)
    h = ParagraphStyle("h", parent=estilos["Heading2"], spaceBefore=12, spaceAfter=6)
    p = ParagraphStyle("p", parent=estilos["BodyText"], leading=15, spaceAfter=6)
    celda = ParagraphStyle("c", parent=estilos["BodyText"], fontSize=9, leading=11)

    def tabla(filas, anchos):
        datos = [[Paragraph(str(c), celda) for c in fila] for fila in filas]
        t = Table(datos, colWidths=anchos, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d8e8d8")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        return t

    e = []
    e.append(Paragraph("Política de Devoluciones y Reembolsos", titulo))
    e.append(Paragraph("EcoMarket · Versión 4.2 · Vigente desde el 1 de agosto de 2026", sub))

    e.append(Paragraph("1. Alcance", h))
    e.append(Paragraph(
        "Esta política aplica a todas las compras realizadas en ecomarket.com. Explica "
        "qué productos pueden devolverse, en qué plazo, cómo se solicita la devolución "
        "y cuándo se devuelve el dinero. Su objetivo es que devolver un producto sea un "
        "proceso sencillo, sin dejar de proteger la salud de nuestros clientes y la "
        "calidad de los productos que reingresan a nuestro inventario.", p))

    e.append(Paragraph("2. Plazo y condiciones generales", h))
    e.append(Paragraph(
        "El cliente tiene 30 días calendario, contados desde la fecha de entrega "
        "registrada por la transportadora, para solicitar la devolución de un producto "
        "retornable por cambio de opinión. Para que la devolución sea aceptada, el "
        "producto debe cumplir estas condiciones:", p))
    for item in [
        "Estar sin usar y en las mismas condiciones en que fue recibido.",
        "Conservar su empaque original y, si lo tiene, su etiqueta o sello.",
        "Incluir todos los accesorios y componentes que venían con el producto.",
    ]:
        e.append(Paragraph("• " + item, p))
    e.append(Paragraph(
        "Si el producto no cumple estas condiciones, EcoMarket puede rechazar la "
        "devolución y se lo comunicará al cliente indicando el motivo.", p))

    e.append(Paragraph("3. Productos que no admiten devolución", h))
    e.append(Paragraph(
        "Por razones de higiene, seguridad o naturaleza del producto, las siguientes "
        "categorías no admiten devolución por cambio de opinión:", p))
    e.append(tabla([
        ["Categoría sin devolución", "Motivo", "Ejemplos"],
        ["Alimentos y bebidas", "Productos perecederos o de consumo", "Granola, café, té de hierbas"],
        ["Higiene personal", "Uso íntimo o personal; no pueden reingresar al inventario por salud pública",
         "Cepillos de dientes, jabones, champú sólido, copa menstrual, pasta dental sólida"],
        ["Productos de limpieza líquidos", "Seguridad: no se reingresan una vez despachados", "Detergente biodegradable"],
        ["Semillas, plantas y seres vivos", "Material biológico de germinación sensible al almacenamiento", "Kits de semillas"],
        ["Ropa interior y trajes de baño", "Prendas de contacto íntimo", "Trajes de baño de plástico reciclado"],
        ["Productos personalizados", "Se fabrican a pedido del cliente", "Cuadernos con nombre grabado"],
        ["Tarjetas de regalo", "No son reembolsables", "Tarjetas de $50.000, $100.000 y $200.000"],
    ], [4.2 * cm, 6.8 * cm, 5.6 * cm]))
    e.append(Spacer(1, 8))
    e.append(Paragraph(
        "La ficha de cada producto en el catálogo indica si es retornable. Cuando un "
        "producto no admite devolución por cambio de opinión, sí se atienden los casos "
        "descritos en la sección 4.", p))

    e.append(Paragraph("4. Productos dañados, defectuosos o enviados por error", h))
    e.append(Paragraph(
        "Estos casos se atienden para todos los productos, incluso los de categorías "
        "que no admiten devolución por cambio de opinión:", p))
    e.append(tabla([
        ["Situación", "Plazo para reportar", "Qué debe enviar el cliente"],
        ["Producto dañado en el transporte, incompleto o equivocado", "5 días calendario desde la entrega",
         "Fotos del paquete y del contenido"],
        ["Defecto de fabricación", "Durante el periodo de garantía del producto (ver documento de garantías)",
         "Fotos o video del defecto"],
        ["Alimento en mal estado o vencido", "48 horas desde la entrega", "Foto del producto y de su fecha de vencimiento"],
    ], [5.6 * cm, 5.4 * cm, 5.6 * cm]))
    e.append(Spacer(1, 8))
    e.append(Paragraph(
        "En estos casos EcoMarket ofrece la reposición del producto o el reembolso "
        "completo, a elección del cliente y sujeto a disponibilidad, y cubre el costo "
        "del envío de retorno.", p))

    e.append(Paragraph("5. Cómo solicitar una devolución", h))
    for n, paso in enumerate([
        "Ingresar a ecomarket.com/devoluciones.",
        "Escribir el número de pedido y el correo electrónico usado en la compra.",
        "Seleccionar el producto a devolver y el motivo.",
        "Descargar e imprimir la etiqueta de envío prepagada.",
        "Entregar el paquete, con la etiqueta pegada, en un punto de recolección de la "
        "transportadora. En Bogotá, Medellín, Cali y Barranquilla también puede solicitarse "
        "la recogida a domicilio.",
    ], 1):
        e.append(Paragraph(f"{n}. {paso}", p))

    e.append(Paragraph("6. Costo del envío de la devolución", h))
    e.append(tabla([
        ["Motivo de la devolución", "Quién asume el envío de retorno"],
        ["Producto defectuoso, dañado o error de EcoMarket", "EcoMarket. La etiqueta es gratuita."],
        ["Cambio de opinión", "El cliente. Se descuentan $8.000 del valor del reembolso."],
    ], [8.3 * cm, 8.3 * cm]))

    e.append(Paragraph("7. Reembolso", h))
    e.append(Paragraph(
        "Cuando el producto llega a nuestra bodega, el equipo de calidad lo inspecciona "
        "en un máximo de 3 días hábiles. Si cumple las condiciones, el reembolso se "
        "procesa en un plazo de 5 a 7 días hábiles y se abona al mismo medio de pago "
        "usado en la compra. Los pagos hechos contra entrega o por transferencia se "
        "reembolsan por transferencia bancaria a la cuenta que indique el cliente.", p))
    e.append(Paragraph(
        "El reembolso incluye el valor del producto. El costo del envío original solo "
        "se reembolsa cuando la devolución se debe a un defecto, un daño o un error de "
        "EcoMarket. El cliente también puede elegir recibir el reembolso como crédito "
        "EcoMarket, que se abona de inmediato con un 5 % adicional.", p))

    e.append(Paragraph("8. Cambios de producto", h))
    e.append(Paragraph(
        "Si el cliente desea cambiar un producto retornable por otro, debe solicitar la "
        "devolución y hacer una compra nueva. Esto garantiza que el producto deseado esté "
        "disponible y que el cliente reciba su reembolso sin demoras.", p))

    e.append(Paragraph("9. Contacto", h))
    e.append(Paragraph(
        "Para dudas sobre devoluciones: soporte@ecomarket.com, chat del sitio web o línea "
        "de atención (601) 555-0100, de lunes a viernes de 8:00 a.m. a 6:00 p.m. y los "
        "sábados de 9:00 a.m. a 1:00 p.m.", p))

    SimpleDocTemplate(
        str(DESTINO / "politica_devoluciones.pdf"), pagesize=A4,
        leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
        title="Política de Devoluciones y Reembolsos - EcoMarket", author="EcoMarket",
    ).build(e)


# ---------------------------------------------------------------------------
# 2. DOCX: garantías y sostenibilidad
# ---------------------------------------------------------------------------
def generar_docx():
    d = Document()
    d.add_heading("Garantías y Compromiso de Sostenibilidad", 0)
    d.add_paragraph("EcoMarket · Versión 2.0 · Vigente desde el 1 de agosto de 2026")

    d.add_heading("1. Qué cubre la garantía", 1)
    d.add_paragraph(
        "EcoMarket respalda sus productos con una garantía por defectos de fabricación, "
        "adicional a los derechos que otorga la Ley 1480 de 2011 (Estatuto del Consumidor). "
        "La garantía cubre fallas de materiales o de construcción que aparezcan con el uso "
        "normal del producto durante el periodo indicado para cada categoría."
    )

    d.add_heading("2. Periodos de garantía", 1)
    d.add_paragraph(
        "El periodo se cuenta desde la fecha de entrega. En el catálogo, la columna "
        "garantia_meses indica el periodo exacto de cada producto."
    )
    filas = [
        ("Producto o material", "Garantía", "Ejemplos"),
        ("Acero inoxidable (termos y botellas)", "24 meses", "Termo de acero 500 ml, Termo familiar 1,2 L"),
        ("Silicona y utensilios portátiles", "12 meses", "Botella plegable, cubiertos portátiles"),
        ("Eco-tecnología solar", "12 meses", "Cargador solar, lámpara solar"),
        ("Compostera y accesorios de hogar duraderos", "12 meses", "Compostera de cocina 5 L, bolsa de lavado para microplásticos"),
        ("Bambú, madera y fibras naturales", "6 meses", "Pajillas de bambú, cepillo de cocina, herramientas de jardín"),
        ("Textiles de algodón orgánico", "3 a 6 meses", "Camiseta (3 meses), bolsas reutilizables (6 meses)"),
        ("Consumibles, papelería, alimentos e higiene", "No aplica", "Cera de abeja, cuadernos, granola, jabones"),
    ]
    t = d.add_table(rows=0, cols=3)
    t.style = "Light Grid Accent 1"
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = v

    d.add_heading("3. Cómo hacer efectiva la garantía", 1)
    for paso in [
        "Escribir a soporte@ecomarket.com o al chat con el número de pedido, fotos o video del defecto y una breve descripción.",
        "El equipo de calidad revisa el caso y responde en un máximo de 5 días hábiles.",
        "Si la garantía procede, EcoMarket repone el producto o reembolsa su valor, a elección del cliente y sujeto a disponibilidad, en un plazo máximo de 30 días.",
        "El envío del producto defectuoso corre por cuenta de EcoMarket.",
    ]:
        d.add_paragraph(paso, style="List Number")

    d.add_heading("4. Qué no cubre la garantía", 1)
    for item in [
        "Golpes, abolladuras, caídas y daños por uso indebido.",
        "Desgaste normal por el uso, como rayones superficiales o decoloración.",
        "Lavado de productos de bambú o madera en lavavajillas, o remojo prolongado.",
        "Daños por exposición a temperaturas extremas fuera de las indicadas en el producto.",
        "Productos modificados, reparados por terceros o usados con fines comerciales.",
    ]:
        d.add_paragraph(item, style="List Bullet")

    d.add_heading("5. Derecho de retracto", 1)
    d.add_paragraph(
        "Para las compras hechas por internet, la ley colombiana reconoce el derecho de "
        "retracto: el cliente puede arrepentirse de la compra dentro de los 5 días hábiles "
        "siguientes a la entrega. EcoMarket va más allá de este mínimo legal y permite "
        "devolver los productos retornables hasta 30 días después de la entrega, según la "
        "Política de Devoluciones y Reembolsos. El derecho de retracto no aplica a "
        "productos que por su naturaleza no admiten devolución, como los perecederos o los "
        "personalizados."
    )

    d.add_heading("6. Programa Recompra de Termos", 1)
    d.add_paragraph(
        "El cliente que entrega un termo o botella usada de acero inoxidable, de cualquier "
        "marca, recibe 15 % de descuento en la compra de un termo nuevo de EcoMarket. El "
        "termo usado se entrega en un punto de recolección de la transportadora o en el "
        "momento de la entrega del pedido nuevo, y el acero se envía a reciclaje. El "
        "descuento no es acumulable con el cupón de bienvenida."
    )

    d.add_heading("7. Compromisos de sostenibilidad", 1)
    for item in [
        "Empaques de cartón reciclado y papel kraft; sin plástico de un solo uso.",
        "Prioridad a proveedores colombianos y a materiales con certificación, como bambú y madera con sello FSC.",
        "Compensación de la huella de carbono del transporte mediante proyectos de reforestación.",
        "Reporte anual de impacto con los kilos de residuos evitados y los productos reciclados.",
    ]:
        d.add_paragraph(item, style="List Bullet")

    d.save(DESTINO / "politica_garantias_sostenibilidad.docx")


# ---------------------------------------------------------------------------
# 3. CSV: catálogo de productos
# ---------------------------------------------------------------------------
CAMPOS = [
    "sku", "nombre", "categoria", "descripcion", "material", "precio_cop", "stock",
    "disponibilidad", "origen", "devolucion_permitida", "motivo_restriccion",
    "garantia_meses", "certificaciones",
]
BODEGA = "En bodega"
PEDIDO = "Bajo pedido (10 a 20 días hábiles)"
AGOTADO = "Agotado"
HIG = "Producto de higiene personal de uso íntimo."
PERE = "Producto perecedero de consumo."

PRODUCTOS = [
    ("PROD-001", "Cepillo de Dientes de Bambú", "Higiene Personal", "Cepillo con mango de bambú y cerdas suaves de nailon biobasado. Presentación individual.", "Bambú Moso", 8900, 120, BODEGA, "Colombia", "No", HIG, 0, "FSC"),
    ("PROD-002", "Termo de Acero Inoxidable 500 ml", "Accesorios Reutilizables", "Termo de doble pared que mantiene bebidas calientes 12 horas y frías 24 horas. Tapa hermética.", "Acero inoxidable 18/8", 69900, 45, BODEGA, "Importado", "Sí", "", 24, "Libre de BPA"),
    ("PROD-003", "Jabón Artesanal Orgánico 100 g", "Higiene Personal", "Jabón en barra de aceite de oliva y avena, sin fragancias sintéticas.", "Aceites vegetales orgánicos", 14500, 80, BODEGA, "Colombia", "No", "Producto perecedero de base orgánica con vida útil limitada.", 0, "Orgánico certificado"),
    ("PROD-004", "Bolsa Reutilizable Pack x3", "Accesorios Reutilizables", "Set de tres bolsas plegables de algodón orgánico para compras, con asas reforzadas.", "Algodón orgánico", 24900, 0, PEDIDO, "Importado", "Sí", "", 6, "Comercio justo"),
    ("PROD-005", "Kit de Pajillas de Bambú x6", "Accesorios Reutilizables", "Seis pajillas de bambú con cepillo de limpieza y estuche de tela.", "Bambú", 19900, 60, BODEGA, "Colombia", "Sí", "", 6, "FSC"),
    ("PROD-006", "Botella Plegable de Silicona 600 ml", "Accesorios Reutilizables", "Botella que se pliega al vaciarse. Apta para congelador y lavavajillas.", "Silicona de grado alimenticio", 39900, 35, BODEGA, "Importado", "Sí", "", 12, "Libre de BPA"),
    ("PROD-007", "Set de Cubiertos Portátiles", "Accesorios Reutilizables", "Tenedor, cuchillo, cuchara y palillos de acero en estuche de corcho.", "Acero inoxidable y corcho", 34900, 28, BODEGA, "Importado", "Sí", "", 12, ""),
    ("PROD-008", "Champú Sólido Natural 80 g", "Higiene Personal", "Barra de champú para todo tipo de cabello, con aceite de coco y romero. Reemplaza dos botellas plásticas.", "Aceites vegetales", 22900, 54, BODEGA, "Colombia", "No", "Producto de higiene personal con vida útil limitada.", 0, "Cruelty free"),
    ("PROD-009", "Envoltura de Cera de Abeja Set x3", "Hogar y Cocina", "Tres envolturas de tela con cera de abeja para conservar alimentos, alternativa al plástico film.", "Algodón, cera de abeja y resina", 28900, 40, BODEGA, "Colombia", "Sí", "", 0, ""),
    ("PROD-010", "Esponja Vegetal de Luffa x2", "Hogar y Cocina", "Dos esponjas naturales de luffa para lavar platos y superficies. Compostables.", "Luffa", 12900, 75, BODEGA, "Colombia", "Sí", "", 0, "Compostable"),
    ("PROD-011", "Cepillo de Cocina de Madera", "Hogar y Cocina", "Cepillo para platos con mango de madera y cerdas vegetales de fibra de agave.", "Madera de haya y agave", 16900, 50, BODEGA, "Colombia", "Sí", "", 6, "FSC"),
    ("PROD-012", "Compostera de Cocina 5 L", "Hogar y Cocina", "Recipiente con filtro de carbón para residuos orgánicos de cocina. Incluye 2 filtros de repuesto.", "Acero y filtro de carbón", 89900, 18, BODEGA, "Importado", "Sí", "", 12, ""),
    ("PROD-013", "Detergente Líquido Biodegradable 1 L", "Hogar y Cocina", "Detergente concentrado para ropa, biodegradable y sin fosfatos.", "Surfactantes de origen vegetal", 31900, 66, BODEGA, "Colombia", "No", "Producto de limpieza líquido; por seguridad no se reingresa al inventario.", 0, "Biodegradable"),
    ("PROD-014", "Granola Orgánica 350 g", "Alimentos y Bebidas", "Granola de avena, almendras y miel orgánica, sin azúcar refinada.", "Avena, almendras y miel", 18900, 60, BODEGA, "Colombia", "No", PERE, 0, "Orgánico certificado"),
    ("PROD-015", "Café de Origen Orgánico 250 g", "Alimentos y Bebidas", "Café molido de cultivo orgánico de pequeños productores del Huila, tostión media.", "Café arábica", 26900, 90, BODEGA, "Colombia", "No", PERE, 0, "Orgánico certificado; Comercio justo"),
    ("PROD-016", "Té de Hierbas Sostenible x20", "Alimentos y Bebidas", "Veinte bolsitas compostables de manzanilla, menta y toronjil.", "Hierbas secas", 15900, 70, BODEGA, "Colombia", "No", PERE, 0, "Compostable"),
    ("PROD-017", "Kit de Semillas de Huerta Urbana", "Jardín y Huerta", "Seis variedades de semillas de hierbas y hortalizas para balcón, con guía de siembra.", "Semillas orgánicas", 21900, 44, BODEGA, "Colombia", "No", "Semillas: material biológico con germinación sensible al almacenamiento.", 0, "Orgánico certificado"),
    ("PROD-018", "Macetas de Fibra de Coco x3", "Jardín y Huerta", "Tres macetas biodegradables que se plantan directamente en la tierra.", "Fibra de coco", 17900, 52, BODEGA, "Colombia", "Sí", "", 6, "Biodegradable"),
    ("PROD-019", "Herramientas de Jardín en Bambú x3", "Jardín y Huerta", "Pala, rastrillo y trasplantador con mangos de bambú y cabezas de acero.", "Bambú y acero", 54900, 22, BODEGA, "Importado", "Sí", "", 6, "FSC"),
    ("PROD-020", "Cuaderno de Papel Reciclado A5", "Papelería y Oficina", "Cuaderno argollado de 100 hojas rayadas, tapa de cartón reciclado.", "Papel 100 % reciclado", 9900, 140, BODEGA, "Colombia", "Sí", "", 0, ""),
    ("PROD-021", "Lápices de Madera Certificada x12", "Papelería y Oficina", "Doce lápices HB de madera con sello FSC y borrador sin PVC.", "Madera FSC", 8500, 200, BODEGA, "Colombia", "Sí", "", 0, "FSC"),
    ("PROD-022", "Agenda 2027 de Papel de Caña", "Papelería y Oficina", "Agenda semanal con tapa dura, papel hecho con bagazo de caña de azúcar.", "Papel de caña", 32900, 0, AGOTADO, "Colombia", "Sí", "", 0, ""),
    ("PROD-023", "Cargador Solar Portátil 10.000 mAh", "Eco-tecnología", "Batería externa con panel solar y dos puertos USB. Resistente a salpicaduras.", "Plástico reciclado y silicio", 129900, 12, BODEGA, "Importado", "Sí", "", 12, ""),
    ("PROD-024", "Lámpara Solar de Jardín x4", "Eco-tecnología", "Cuatro lámparas LED con panel solar y sensor de luz para exteriores.", "Plástico reciclado y silicio", 74900, 0, AGOTADO, "Importado", "Sí", "", 12, ""),
    ("PROD-025", "Camiseta de Algodón Orgánico", "Ropa y Textiles", "Camiseta unisex de corte clásico disponible en tallas S, M, L y XL.", "Algodón orgánico", 59900, 64, BODEGA, "Colombia", "Sí", "", 3, "Orgánico certificado"),
    ("PROD-026", "Bolsa de Lavado para Microplásticos", "Ropa y Textiles", "Bolsa para la lavadora que retiene las fibras sintéticas antes de que lleguen al agua.", "Poliamida reciclada", 44900, 31, BODEGA, "Importado", "Sí", "", 12, ""),
    ("PROD-027", "Traje de Baño de Plástico Reciclado", "Ropa y Textiles", "Traje de baño de una pieza hecho con redes de pesca recuperadas, tallas S a XL.", "Nailon regenerado", 89900, 20, BODEGA, "Importado", "No", "Prenda de contacto íntimo; no se acepta devolución por razones de higiene.", 0, ""),
    ("PROD-028", "Copa Menstrual de Silicona Médica", "Higiene Personal", "Copa reutilizable con duración de hasta 10 años, disponible en tallas A y B.", "Silicona médica", 49900, 38, BODEGA, "Importado", "No", HIG, 0, "Libre de BPA"),
    ("PROD-029", "Pasta Dental Sólida en Tabletas x60", "Higiene Personal", "Sesenta tabletas masticables con flúor, en frasco de vidrio reutilizable.", "Bicarbonato y xilitol", 27900, 58, BODEGA, "Colombia", "No", HIG, 0, "Cruelty free"),
    ("PROD-030", "Termo Familiar de Acero 1,2 L", "Accesorios Reutilizables", "Termo de gran capacidad con asa y tapa tipo vaso. Mantiene el calor 18 horas.", "Acero inoxidable 18/8", 99900, 9, BODEGA, "Importado", "Sí", "", 24, "Libre de BPA"),
]


def generar_csv():
    with open(DESTINO / "catalogo_productos.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(CAMPOS)
        w.writerows(PRODUCTOS)


if __name__ == "__main__":
    generar_pdf()
    generar_docx()
    generar_csv()
    print("Documentos generados en", DESTINO)
